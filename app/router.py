"""
Decide que motor usar segun el proyecto y la forma del archivo, y ejecuta
el procesamiento completo. Aisla esta logica de main.py para que agregar
un proyecto multi-codigo nuevo (como HCA) sea: escribir su perfil en
profiles/, y registrarlo aqui -- nada mas.
"""

import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.process_simple import (
    process as process_simple, write_result as write_simple, add_missing_employees, sync_rate_store,
    roster_position_breakdown, write_invoice_totals,
)
from engine import rate_store as rate_store_mod
from engine.process_auto import build_profile_auto
from engine.process_daisy import read_daisy_source, write_daisy_result
from engine.matching import match_source_to_roster
from engine import process_hca

# Proyectos "multi-codigo" conocidos: nombre normalizado (minusculas, sin
# espacios extra) -> modulo de perfil. Un proyecto nuevo de este tipo se
# agrega aqui una vez definido su mapeo fila->codigo.
_MULTICODE_PROJECTS = {}


def _register_multicode(name: str, profile, tier_fn):
    _MULTICODE_PROJECTS[name.strip().lower()] = (profile, tier_fn)


def _load_multicode_registry():
    if _MULTICODE_PROJECTS:
        return
    from profiles.hca_palms_west import PROFILE as hca_profile, tier_fn as hca_tier_fn
    _register_multicode(hca_profile.project_name, hca_profile, hca_tier_fn)


def _looks_like_daisy_source(ws) -> bool:
    """El formato Project Daisy se distingue por traer 'RT Hours' / 'Cost
    Code #' como encabezados, en vez de columnas de dia (MON, TUE, ...)."""
    for r in range(1, 15):
        for c in range(1, 25):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper() in ("RT HOURS", "COST CODE #"):
                return True
    return False


class ProcessOutcome:
    def __init__(self, pipeline, matched_count, missing_people, ambiguous, breakdown, totals=None, error=None):
        self.pipeline = pipeline
        self.matched_count = matched_count
        self.missing_people = missing_people
        self.ambiguous = ambiguous
        self.breakdown = breakdown
        self.totals = totals or {}
        self.error = error


def _ensure_xlsx(path: Path) -> Path:
    """Convierte .xls (formato viejo de Excel) a .xlsx con LibreOffice
    headless antes de procesarlo -- necesario para fuentes tipo Project
    Daisy que a veces llegan en ese formato."""
    if path.suffix.lower() != ".xls":
        return path
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(path.parent), str(path)],
        check=True, capture_output=True, timeout=60,
    )
    return path.with_suffix(".xlsx")


def process_upload(project_name: str, source_path: Path, template_path: Path, output_path: Path) -> ProcessOutcome:
    _load_multicode_registry()
    key = project_name.strip().lower()
    source_path = _ensure_xlsx(source_path)

    if key in _MULTICODE_PROJECTS:
        return _process_multicode(key, source_path, template_path, output_path)

    source_wb = openpyxl.load_workbook(source_path, data_only=True)
    source_ws = source_wb["Sheet1"] if "Sheet1" in source_wb.sheetnames else source_wb.active

    if _looks_like_daisy_source(source_ws):
        return _process_daisy(project_name, source_ws, template_path, output_path)

    dest_wb = openpyxl.load_workbook(template_path)
    dest_ws = dest_wb["PH"] if "PH" in dest_wb.sheetnames else dest_wb.active
    return _process_standard(project_name, source_ws, dest_wb, dest_ws, output_path)


def _process_standard(project_name, source_ws, dest_wb, dest_ws, output_path):
    try:
        profile = build_profile_auto(project_name, source_ws, dest_ws, dest_sheet_name=dest_ws.title, source_sheet_name=source_ws.title)
    except ValueError as e:
        return ProcessOutcome("standard", 0, [], [], [], error=str(e))

    result = process_simple(source_ws, dest_ws, profile)
    write_simple(dest_ws, profile, result)

    store = rate_store_mod.load(project_name)
    gstore = rate_store_mod.load_global()
    added = add_missing_employees(dest_ws, profile, result.missing_people, rate_store=store, global_store=gstore)
    sync_rate_store(dest_ws, profile, result.filled_rows, added, store, gstore)
    rate_store_mod.save(project_name, store)
    rate_store_mod.save_global(gstore)

    # horas reales por posicion, ya con los agregados incluidos -- se usa
    # para el resumen y, si la plantilla trae el cuadro de "invoice"
    # (horas x tarifa debajo del roster), para dejarlo resuelto/verificado
    # en vez de un numero viejo copiado a mano.
    real_breakdown = roster_position_breakdown(dest_ws, profile)
    invoice_updated = write_invoice_totals(dest_ws, profile, real_breakdown)

    dest_wb.save(output_path)

    breakdown = sorted(real_breakdown.items())
    known_added = [a for a in added if a[2]]
    new_added = [a for a in added if not a[2]]
    note = []
    if invoice_updated:
        note.append("Se actualizo el cuadro de INVOICE (horas por posicion) al final de la hoja con las horas reales de esta semana -- revisalo antes de facturar.")
    if known_added:
        names = ", ".join(e.name for _, e, _ in known_added)
        note.append(f"{len(known_added)} persona(s) que ya conociamos (de este u otro proyecto), agregada(s) con su INTAL_RATE real: {names}.")
    if new_added:
        names = ", ".join(e.name for _, e, _ in new_added)
        note.append(
            f"{len(new_added)} persona(s) totalmente nueva(s) agregada(s) en amarillo con INTAL_RATE y RATE en 0: {names} "
            "-- revisar y completar su tarifa a mano."
        )
    outcome = ProcessOutcome("standard", len(result.filled_rows), [], result.ambiguous, breakdown)
    outcome.totals["notes"] = note
    return outcome


def _process_daisy(project_name, source_ws, template_path, output_path):
    dest_wb = openpyxl.load_workbook(template_path)
    cobrar_sheet = next((s for s in dest_wb.sheetnames if "COBRAR" in s.upper()), dest_wb.sheetnames[0])
    dest_ws = dest_wb[cobrar_sheet]

    employees = read_daisy_source(source_ws)
    day_employees = [e for e in employees if e.shift == "DAY"]
    night_employees = [e for e in employees if e.shift == "NIGHT"]

    roster = {}
    for r in range(10, dest_ws.max_row + 1):
        name = dest_ws.cell(r, 2).value
        if name and str(name).strip():
            roster[r] = str(name).strip()

    match = match_source_to_roster([e.name for e in day_employees], roster)
    write_daisy_result(dest_ws, day_employees, match.matches)
    dest_wb.save(output_path)

    missing = [type("M", (), {"name": e.name, "skill": "", "days": e.days, "source_reg": e.reg, "source_ot": e.ot}) for e in day_employees if e.name in match.unmatched_source]

    note = []
    if night_employees:
        note.append(f"{len(night_employees)} combinaciones de turno NOCHE detectadas -- van al bloque de totales por posicion, revisar a mano esta primera vez")

    breakdown = {}
    for e in day_employees + night_employees:
        bucket = e.shift
        d = breakdown.setdefault(bucket, {"reg": 0.0, "ot": 0.0})
        d["reg"] += e.reg
        d["ot"] += e.ot

    outcome = ProcessOutcome("daisy", len(match.matches), missing, match.ambiguous, sorted(breakdown.items()))
    outcome.totals["notes"] = note
    return outcome


def _process_multicode(key, source_path, template_path, output_path):
    profile, tier_fn = _MULTICODE_PROJECTS[key]

    source_wb = openpyxl.load_workbook(source_path, data_only=True)
    source_ws = source_wb[profile.source_sheet]

    dest_wb = openpyxl.load_workbook(template_path)
    ts_ws = dest_wb[profile.dest_timesheet_sheet]
    cb_ws = dest_wb[profile.dest_cobrar_sheet] if profile.dest_cobrar_sheet else None

    recs = process_hca.read_source(source_ws, profile)

    # detectar combinaciones nuevas (persona+codigo) que no estan en el mapeo guardado
    known_keys = set(profile.row_to_source_key.values())
    new_keys = [k for k in recs.keys() if k not in known_keys]

    regot = process_hca.compute_regot(profile, recs)
    process_hca.write_timesheet(ts_ws, profile, recs, regot)

    error = None
    if cb_ws is not None:
        try:
            slots = process_hca.detect_cobrar_slots(cb_ws)
            groups = process_hca.build_cobrar_groups(ts_ws, profile, regot, tier_fn)
            process_hca.write_cobrar(cb_ws, slots, groups)
        except ValueError as e:
            error = str(e)

    dest_wb.save(output_path)

    missing = [
        type("M", (), {"name": f"{n} (codigo {c}, nunca definido)", "skill": "", "days": recs[(n, c)]["days"], "source_reg": recs[(n, c)]["rt"], "source_ot": recs[(n, c)]["ot"]})
        for (n, c) in new_keys
    ]

    breakdown = {}
    for row, (reg, ot) in regot.items():
        pos = ts_ws.cell(row, 3).value or "SIN CLASIFICAR"
        d = breakdown.setdefault(pos.strip(), {"reg": 0.0, "ot": 0.0})
        d["reg"] += reg
        d["ot"] += ot

    return ProcessOutcome("multicode", len(regot), missing, [], sorted(breakdown.items()), error=error)
